# app.py
import streamlit as st
import pandas as pd
import PyPDF2
import openpyxl
import xlrd
import requests
import json
import re
import io
import os
import traceback
from typing import Dict, List, Optional, Tuple, Any
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from datetime import datetime
import nltk
import textstat

# Download required NLTK data
try:
    nltk.data.find('tokenizers/punkt')
except LookupError:
    nltk.download('punkt', quiet=True)

try:
    nltk.data.find('corpora/stopwords')
except LookupError:
    nltk.download('stopwords', quiet=True)

from nltk.corpus import stopwords
from nltk.tokenize import word_tokenize, sent_tokenize

class FinancialDocumentProcessor:
    """Main class for processing financial documents and handling Q&A"""
    
    def __init__(self):
        self.financial_keywords = {
            'revenue': ['revenue', 'sales', 'income', 'turnover', 'gross sales'],
            'expenses': ['expenses', 'costs', 'expenditure', 'outgoing', 'spending'],
            'profit': ['profit', 'earnings', 'net income', 'surplus', 'margin'],
            'assets': ['assets', 'holdings', 'property', 'investments'],
            'liabilities': ['liabilities', 'debt', 'obligations', 'payables'],
            'equity': ['equity', 'shareholders equity', 'owners equity', 'capital'],
            'cash': ['cash', 'cash flow', 'liquidity', 'cash equivalents'],
            'ratios': ['ratio', 'margin', 'percentage', 'rate', 'return']
        }
        
        self.document_data = {}
        self.extracted_text = ""
        self.financial_metrics = {}
        self.tables = []
        
    def extract_pdf_content(self, pdf_file) -> Tuple[str, bool]:
        """Extract text content from PDF file"""
        try:
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            text = ""
            
            for page_num, page in enumerate(pdf_reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n--- Page {page_num + 1} ---\n{page_text}"
                except Exception as e:
                    st.warning(f"Could not extract text from page {page_num + 1}: {str(e)}")
                    continue
            
            if not text.strip():
                return "No text could be extracted from the PDF.", False
                
            return text, True
            
        except Exception as e:
            return f"Error processing PDF: {str(e)}", False
    
    def extract_excel_content(self, excel_file) -> Tuple[str, bool]:
        """Extract content from Excel file"""
        try:
            # Try different engines
            dfs = {}
            text_content = ""
            
            try:
                # Try openpyxl first (for .xlsx files)
                xl_file = pd.ExcelFile(excel_file, engine='openpyxl')
            except:
                try:
                    # Try xlrd for .xls files
                    xl_file = pd.ExcelFile(excel_file, engine='xlrd')
                except Exception as e:
                    return f"Could not read Excel file: {str(e)}", False
            
            for sheet_name in xl_file.sheet_names:
                try:
                    df = pd.read_excel(xl_file, sheet_name=sheet_name)
                    dfs[sheet_name] = df
                    
                    # Convert dataframe to text representation
                    text_content += f"\n--- Sheet: {sheet_name} ---\n"
                    text_content += df.to_string(index=False, na_rep='')
                    text_content += "\n"
                    
                    # Store tables for later use
                    self.tables.append({
                        'sheet_name': sheet_name,
                        'data': df,
                        'summary': f"Sheet '{sheet_name}' contains {df.shape[0]} rows and {df.shape[1]} columns"
                    })
                    
                except Exception as e:
                    st.warning(f"Could not process sheet '{sheet_name}': {str(e)}")
                    continue
            
            if not text_content.strip():
                return "No data could be extracted from the Excel file.", False
                
            return text_content, True
            
        except Exception as e:
            return f"Error processing Excel file: {str(e)}", False
    
    def extract_financial_metrics(self, text: str) -> Dict[str, Any]:
        """Extract financial metrics from text using pattern matching"""
        metrics = {}
        
        try:
            # Common financial patterns
            patterns = {
                'currency_amounts': r'[\$£€¥₹]\s*[\d,]+\.?\d*[KMB]?|\b\d{1,3}(?:,\d{3})*(?:\.\d{2})?\s*(?:million|billion|thousand|M|B|K)?\b',
                'percentages': r'\d+\.?\d*\s*%',
                'financial_terms': r'\b(?:revenue|sales|profit|loss|assets|liabilities|equity|cash|expenses|costs|income|earnings|margin|ratio|ROI|ROE|EBITDA)\b',
                'dates': r'\b(?:Q[1-4]|FY|fiscal year|\d{4}|\d{1,2}/\d{1,2}/\d{4}|\d{1,2}-\d{1,2}-\d{4}|January|February|March|April|May|June|July|August|September|October|November|December)\b'
            }
            
            for pattern_name, pattern in patterns.items():
                matches = re.findall(pattern, text, re.IGNORECASE)
                if matches:
                    metrics[pattern_name] = list(set(matches))  # Remove duplicates
            
            # Try to identify key financial figures
            lines = text.split('\n')
            financial_data = {}
            
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                    
                # Look for financial statement items
                for category, keywords in self.financial_keywords.items():
                    for keyword in keywords:
                        if keyword.lower() in line.lower():
                            # Extract numbers from the line
                            numbers = re.findall(r'[\d,]+\.?\d*', line)
                            if numbers:
                                if category not in financial_data:
                                    financial_data[category] = []
                                financial_data[category].extend(numbers)
            
            metrics['extracted_financial_data'] = financial_data
            
        except Exception as e:
            st.error(f"Error extracting financial metrics: {str(e)}")
            
        return metrics
    
    def process_document(self, uploaded_file) -> Tuple[str, bool, Dict]:
        """Process uploaded document and extract content"""
        try:
            file_extension = uploaded_file.name.split('.')[-1].lower()
            
            if file_extension == 'pdf':
                text, success = self.extract_pdf_content(uploaded_file)
            elif file_extension in ['xlsx', 'xls']:
                text, success = self.extract_excel_content(uploaded_file)
            else:
                return f"Unsupported file format: {file_extension}", False, {}
            
            if success:
                self.extracted_text = text
                self.financial_metrics = self.extract_financial_metrics(text)
                
                # Store document metadata
                self.document_data = {
                    'filename': uploaded_file.name,
                    'file_type': file_extension,
                    'processed_at': datetime.now().isoformat(),
                    'text_length': len(text),
                    'word_count': len(text.split()),
                    'has_tables': len(self.tables) > 0
                }
                
            return text, success, self.financial_metrics
            
        except Exception as e:
            error_msg = f"Error processing document: {str(e)}\n{traceback.format_exc()}"
            return error_msg, False, {}

class OllamaInterface:
    """Interface for communicating with Ollama local LLM"""
    
    def __init__(self, base_url: str = "http://localhost:11434"):
        self.base_url = base_url
        self.model = "llama2"  # Default model
        
    def check_ollama_connection(self) -> Tuple[bool, str]:
        """Check if Ollama is running and accessible"""
        try:
            response = requests.get(f"{self.base_url}/api/tags", timeout=5)
            if response.status_code == 200:
                models = response.json().get('models', [])
                if models:
                    return True, f"Connected successfully. Available models: {[m['name'] for m in models]}"
                else:
                    return False, "Ollama is running but no models are installed."
            else:
                return False, f"Ollama responded with status code: {response.status_code}"
        except requests.exceptions.ConnectionError:
            return False, "Cannot connect to Ollama. Please make sure Ollama is running on localhost:11434"
        except requests.exceptions.Timeout:
            return False, "Connection to Ollama timed out."
        except Exception as e:
            return False, f"Error connecting to Ollama: {str(e)}"
    
    def get_available_models(self) -> List[str]:
        """Get list of available models from Ollama"""
        try:
            response = requests.get(f"{self.base_url}/api/tags", timeout=5)
            if response.status_code == 200:
                models = response.json().get('models', [])
                return [model['name'] for model in models]
            return []
        except:
            return []
    
    def generate_response(self, prompt: str, context: str = "", model: str = None) -> Tuple[str, bool]:
        """Generate response using Ollama"""
        try:
            if model:
                self.model = model
            
            # Construct the full prompt with context
            full_prompt = f"""You are a financial document analysis assistant. You have access to financial document content and should answer questions accurately based on this information.

Document Content:
{context[:4000]}  # Limit context to avoid token limits

User Question: {prompt}

Please provide a clear, accurate answer based on the financial document content. If you cannot find specific information in the document, please state that clearly. Focus on financial metrics, trends, and insights that can be derived from the provided data.

Answer:"""

            payload = {
                "model": self.model,
                "prompt": full_prompt,
                "stream": False,
                "options": {
                    "temperature": 0.3,  # Lower temperature for more factual responses
                    "num_ctx": 4096,     # Context window
                    "top_p": 0.9
                }
            }
            
            response = requests.post(
                f"{self.base_url}/api/generate",
                json=payload,
                timeout=60
            )
            
            if response.status_code == 200:
                result = response.json()
                return result.get('response', 'No response generated'), True
            else:
                return f"Error: Received status code {response.status_code}", False
                
        except requests.exceptions.Timeout:
            return "Request timed out. The model might be processing a complex query.", False
        except Exception as e:
            return f"Error generating response: {str(e)}", False

def create_financial_dashboard(metrics: Dict, tables: List) -> None:
    """Create a financial dashboard with visualizations"""
    try:
        st.subheader("📊 Financial Data Dashboard")
        
        # Display extracted financial data
        if 'extracted_financial_data' in metrics and metrics['extracted_financial_data']:
            col1, col2 = st.columns(2)
            
            with col1:
                st.write("**Financial Categories Found:**")
                for category, values in metrics['extracted_financial_data'].items():
                    if values:
                        st.write(f"• {category.title()}: {len(values)} entries")
            
            with col2:
                st.write("**Document Analysis:**")
                if 'currency_amounts' in metrics:
                    st.write(f"• Currency amounts found: {len(metrics['currency_amounts'])}")
                if 'percentages' in metrics:
                    st.write(f"• Percentages found: {len(metrics['percentages'])}")
        
        # Display tables if available
        if tables:
            st.subheader("📋 Data Tables")
            for table in tables:
                with st.expander(f"View {table['sheet_name']} ({table['summary']})"):
                    st.dataframe(table['data'])
                    
                    # Try to create simple visualizations for numeric data
                    numeric_columns = table['data'].select_dtypes(include=[np.number]).columns
                    if len(numeric_columns) > 0:
                        st.write("**Quick Visualization:**")
                        try:
                            # Simple bar chart of numeric columns
                            if len(numeric_columns) <= 5:  # Limit to avoid cluttered charts
                                fig = px.bar(
                                    x=numeric_columns.tolist(),
                                    y=[table['data'][col].sum() if not table['data'][col].isnull().all() else 0 for col in numeric_columns],
                                    title=f"Sum of Numeric Columns in {table['sheet_name']}"
                                )
                                st.plotly_chart(fig, use_container_width=True)
                        except Exception as e:
                            st.write(f"Could not create visualization: {str(e)}")
        
        # Display extracted metrics
        if metrics:
            st.subheader("🔍 Extracted Information")
            for key, value in metrics.items():
                if key != 'extracted_financial_data' and value:
                    with st.expander(f"View {key.replace('_', ' ').title()}"):
                        if isinstance(value, list):
                            for item in value[:20]:  # Limit display to first 20 items
                                st.write(f"• {item}")
                        else:
                            st.write(value)
                            
    except Exception as e:
        st.error(f"Error creating dashboard: {str(e)}")

def main():
    """Main Streamlit application"""
    st.set_page_config(
        page_title="Financial Document Q&A Assistant",
        page_icon="💼",
        layout="wide",
        initial_sidebar_state="expanded"
    )
    
    # Initialize session state
    if 'processor' not in st.session_state:
        st.session_state.processor = FinancialDocumentProcessor()
    
    if 'ollama' not in st.session_state:
        st.session_state.ollama = OllamaInterface()
    
    if 'chat_history' not in st.session_state:
        st.session_state.chat_history = []
    
    if 'document_processed' not in st.session_state:
        st.session_state.document_processed = False
    
    # Header
    st.title("💼 Financial Document Q&A Assistant")
    st.markdown("Upload financial documents (PDF/Excel) and ask questions about your financial data using natural language.")
    
    # Sidebar
    with st.sidebar:
        st.header("🔧 Configuration")
        
        # Ollama connection check
        st.subheader("Ollama Connection")
        if st.button("Check Ollama Connection"):
            connected, message = st.session_state.ollama.check_ollama_connection()
            if connected:
                st.success(message)
                # Get available models
                models = st.session_state.ollama.get_available_models()
                if models and len(models) > 1:
                    selected_model = st.selectbox("Select Model", models)
                    st.session_state.ollama.model = selected_model
            else:
                st.error(message)
                st.info("To set up Ollama:\n1. Install Ollama from https://ollama.ai\n2. Run `ollama pull llama2`\n3. Start Ollama service")
        
        st.divider()
        
        # Document upload section
        st.subheader("📁 Document Upload")
        uploaded_file = st.file_uploader(
            "Choose a financial document",
            type=['pdf', 'xlsx', 'xls'],
            help="Upload PDF financial statements or Excel spreadsheets"
        )
        
        if uploaded_file is not None:
            if st.button("Process Document", type="primary"):
                with st.spinner("Processing document..."):
                    try:
                        text, success, metrics = st.session_state.processor.process_document(uploaded_file)
                        
                        if success:
                            st.success("✅ Document processed successfully!")
                            st.session_state.document_processed = True
                            
                            # Display document info
                            doc_info = st.session_state.processor.document_data
                            st.info(f"""
                            📄 **Document Info:**
                            - Filename: {doc_info['filename']}
                            - Type: {doc_info['file_type'].upper()}
                            - Words: {doc_info['word_count']:,}
                            - Characters: {doc_info['text_length']:,}
                            - Tables: {'Yes' if doc_info['has_tables'] else 'No'}
                            """)
                        else:
                            st.error(f"❌ Error processing document: {text}")
                            st.session_state.document_processed = False
                            
                    except Exception as e:
                        st.error(f"❌ Unexpected error: {str(e)}")
                        st.session_state.document_processed = False
        
        # Clear chat history
        if st.button("Clear Chat History"):
            st.session_state.chat_history = []
            st.rerun()
    
    # Main content area
    if st.session_state.document_processed:
        # Create tabs for different views
        tab1, tab2, tab3 = st.tabs(["💬 Q&A Chat", "📊 Dashboard", "📄 Document Content"])
        
        with tab1:
            st.header("Ask Questions About Your Financial Document")
            
            # Display chat history
            if st.session_state.chat_history:
                st.subheader("Chat History")
                for i, (question, answer) in enumerate(st.session_state.chat_history):
                    with st.container():
                        st.markdown(f"**You:** {question}")
                        st.markdown(f"**Assistant:** {answer}")
                        st.divider()
            
            # Question input
            col1, col2 = st.columns([4, 1])
            
            with col1:
                user_question = st.text_input(
                    "Ask a question about your financial document:",
                    placeholder="e.g., What was the total revenue for last quarter?",
                    key="question_input"
                )
            
            with col2:
                ask_button = st.button("Ask", type="primary", use_container_width=True)
            
            # Predefined question buttons
            st.subheader("Quick Questions")
            col1, col2, col3, col4 = st.columns(4)
            
            quick_questions = [
                "What are the main revenue sources?",
                "What were the total expenses?",
                "Show me the profit margins",
                "What are the key financial ratios?"
            ]
            
            for i, question in enumerate(quick_questions):
                col = [col1, col2, col3, col4][i]
                if col.button(question, key=f"quick_{i}"):
                    user_question = question
                    ask_button = True
            
            # Process question
            if ask_button and user_question:
                if user_question.strip():
                    with st.spinner("Generating response..."):
                        try:
                            # Check Ollama connection first
                            connected, _ = st.session_state.ollama.check_ollama_connection()
                            
                            if connected:
                                response, success = st.session_state.ollama.generate_response(
                                    user_question,
                                    st.session_state.processor.extracted_text
                                )
                                
                                if success:
                                    st.session_state.chat_history.append((user_question, response))
                                    st.rerun()
                                else:
                                    st.error(f"Error generating response: {response}")
                            else:
                                st.error("Please check Ollama connection first.")
                                
                        except Exception as e:
                            st.error(f"Error processing question: {str(e)}")
                else:
                    st.warning("Please enter a question.")
        
        with tab2:
            create_financial_dashboard(
                st.session_state.processor.financial_metrics,
                st.session_state.processor.tables
            )
        
        with tab3:
            st.header("Document Content Preview")
            
            # Show document statistics
            col1, col2, col3, col4 = st.columns(4)
            doc_data = st.session_state.processor.document_data
            
            col1.metric("Word Count", f"{doc_data['word_count']:,}")
            col2.metric("Characters", f"{doc_data['text_length']:,}")
            col3.metric("File Type", doc_data['file_type'].upper())
            col4.metric("Tables Found", len(st.session_state.processor.tables))
            
            # Show extracted text (truncated)
            st.subheader("Extracted Text (First 2000 characters)")
            extracted_text = st.session_state.processor.extracted_text
            if extracted_text:
                st.text_area(
                    "Document Content:",
                    value=extracted_text[:2000] + ("..." if len(extracted_text) > 2000 else ""),
                    height=400,
                    disabled=True
                )
            
            # Download processed content
            if st.download_button(
                label="Download Processed Content",
                data=extracted_text,
                file_name=f"processed_{doc_data['filename']}.txt",
                mime="text/plain"
            ):
                st.success("Content downloaded successfully!")
    
    else:
        # Welcome screen
        st.header("Welcome to Financial Document Q&A Assistant! 🚀")
        
        col1, col2 = st.columns([2, 1])
        
        with col1:
            st.markdown("""
            ### How to use this application:
            
            1. **Setup Ollama**: Make sure Ollama is installed and running locally
            2. **Upload Document**: Choose a PDF or Excel financial document
            3. **Process**: Click 'Process Document' to extract and analyze content
            4. **Ask Questions**: Use natural language to query your financial data
            5. **Explore**: View dashboards and document content in different tabs
            
            ### Supported Document Types:
            - 📄 **PDF Files**: Income statements, balance sheets, cash flow statements
            - 📊 **Excel Files**: Financial spreadsheets, budgets, financial models
            
            ### Sample Questions You Can Ask:
            - "What was the total revenue for the last quarter?"
            - "Show me the expense breakdown"
            - "What are the profit margins?"
            - "Calculate the debt-to-equity ratio"
            - "What trends do you see in the financial data?"
            """)
        
        with col2:
            st.info("""
            **System Requirements:**
            - Ollama installed locally
            - Python 3.8+
            - Supported models:
              - llama2
              - mistral
              - codellama
            """)
            
            st.success("""
            **Features:**
            ✅ PDF & Excel processing
            ✅ Natural language Q&A
            ✅ Financial metrics extraction
            ✅ Interactive dashboards
            ✅ Chat history
            ✅ Local deployment
            """)

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        st.error(f"Application error: {str(e)}")
        st.error("Please check your setup and try again.")
        st.info("If problems persist, check the console for detailed error messages.")